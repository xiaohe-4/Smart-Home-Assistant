# -*- coding: utf-8 -*-
"""按《材料要求》生成海报、PPT、开发日志 PDF、证明材料 ZIP。"""
from __future__ import annotations

import json
import shutil
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
MAT = Path(__file__).resolve().parent
PROOF = MAT / "07_其他证明材料"
FONT_PATH = Path(r"C:\Windows\Fonts\msyh.ttc")
SIMHEI = Path(r"C:\Windows\Fonts\simhei.ttf")

GREEN = RGBColor(0x16, 0xA3, 0x4A)
DARK = RGBColor(0x08, 0x24, 0x1C)
CREAM = RGBColor(0xF4, 0xFA, 0xF6)
TEXT = RGBColor(0x12, 0x30, 0x26)
MUTED = RGBColor(0x5F, 0x74, 0x6A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD = RGBColor(0xF5, 0x9E, 0x0B)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_run_font(run, name="微软雅黑", size=18, bold=False, color=TEXT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    r = run._r
    rPr = r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


def add_rect(slide, l, t, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_text(slide, l, t, w, h, text, size=18, bold=False, color=TEXT, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, size=size, bold=bold, color=color)
    return box


def bullets(slide, l, t, w, h, lines, size=18):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = "•  " + line
        set_run_font(run, size=size, color=TEXT)
    return box


def footer(slide, page, total=18):
    add_text(slide, Inches(0.5), Inches(7.15), Inches(8), Inches(0.3), "智家能  ·  产品说明", size=11, color=MUTED)
    add_text(slide, Inches(11.2), Inches(7.15), Inches(1.6), Inches(0.3), f"{page} / {total}", size=11, color=MUTED, align=PP_ALIGN.RIGHT)


def new_light(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, CREAM)
    add_rect(slide, 0, 0, Inches(0.12), SLIDE_H, GREEN)
    return slide


def make_poster():
    src_candidates = [
        Path(r"C:\Users\yumen\.cursor\projects\c-Users-yumen-Desktop\assets\poster_source.png"),
        MAT / "poster_source.png",
    ]
    out = MAT / "02_产品海报_智家能_1200x560.png"
    src = next((p for p in src_candidates if p.exists()), None)
    if src:
        im = Image.open(src).convert("RGB")
        tw, th = 1200, 560
        ratio = tw / th
        w, h = im.size
        if w / h > ratio:
            nw = int(h * ratio)
            left = (w - nw) // 2
            im = im.crop((left, 0, left + nw, h))
        else:
            nh = int(w / ratio)
            top = (h - nh) // 2
            im = im.crop((0, top, w, top + nh))
        im = im.resize((tw, th), Image.Resampling.LANCZOS)
        im.save(out, "PNG", optimize=True)
        return out
    # fallback
    img = Image.new("RGB", (1200, 560), (8, 36, 28))
    draw = ImageDraw.Draw(img)
    font_big = ImageFont.truetype(str(FONT_PATH), 92, index=0)
    font_mid = ImageFont.truetype(str(FONT_PATH), 28, index=0)
    font_s = ImageFont.truetype(str(FONT_PATH), 22, index=0)
    draw.rectangle((0, 0, 16, 560), fill=(22, 163, 74))
    draw.text((56, 150), "智家能", font=font_big, fill=(255, 255, 255))
    draw.text((60, 270), "多模态家庭能源智能决策", font=font_mid, fill=(187, 247, 208))
    draw.text((60, 340), "看懂用电    发现浪费    一句话去办", font=font_s, fill=(244, 250, 246))
    draw.text((60, 470), "XGBoost 预测  ·  异常可解释  ·  管家可执行", font=font_s, fill=(134, 239, 172))
    img.save(out, "PNG", optimize=True)
    return out


def make_pptx(poster: Path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 1 cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, DARK)
    add_rect(s, 0, 0, Inches(0.18), SLIDE_H, GREEN)
    add_text(s, Inches(0.8), Inches(1.6), Inches(11), Inches(1.2), "智家能", size=60, bold=True, color=WHITE)
    add_text(s, Inches(0.8), Inches(2.9), Inches(11), Inches(0.6), "基于多模态家庭数据分析的 AI 能源管理与节能决策系统", size=22, color=RGBColor(0xBB, 0xF7, 0xD0))
    add_text(s, Inches(0.8), Inches(4.0), Inches(11), Inches(0.4), "看懂用电  ·  发现浪费  ·  一句话去办", size=20, color=GOLD)
    add_text(s, Inches(0.8), Inches(5.6), Inches(11), Inches(0.4), "演示家庭：成都四口之家  ·  2026-08-19 20:30  ·  已运行 30 天", size=16, color=RGBColor(0x9E, 0xCC, 0xB8))

    # 2 pain
    s = new_light(prs)
    add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "项目背景 / 用户痛点", size=28, bold=True, color=DARK)
    bullets(s, Inches(0.7), Inches(1.2), Inches(12), Inches(5.2), [
        "家里只有一个总电量，看不懂为什么贵。",
        "高峰、平时、谷时混在一起，浪费看不见。",
        "凌晨热水器异常启动，没有规则说明。",
        "建议停留在「请节约用电」，没有人去拧开关。",
        "评委和用户真正要的是：已经跑起来的决策闭环。",
    ], size=22)
    footer(s, 2)

    # 3 positioning
    s = new_light(prs)
    add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "产品定位", size=28, bold=True, color=DARK)
    bullets(s, Inches(0.7), Inches(1.2), Inches(12), Inches(5.2), [
        "面向普通家庭的多模态能源智能决策系统。",
        "不是单纯电量看板，也不是必须上硬件才能演示的中控。",
        "回答三件事：为什么耗电、哪里浪费、今晚谁去办。",
        "多模态 = 电力 / 室内环境 / 占用行为 / 天气 四路时序，不是图像。",
        "大模型只解释，不算指标。",
    ], size=22)
    footer(s, 3)

    # 4 scheme
    s = new_light(prs)
    add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "产品整体方案", size=28, bold=True, color=DARK)
    steps = ["四路感知", "融合对齐", "预测诊断", "多目标决策", "管家执行", "生活反馈"]
    for i, name in enumerate(steps):
        left = Inches(0.55 + i * 2.1)
        add_rect(s, left, Inches(2.4), Inches(1.9), Inches(1.5), DARK)
        add_text(s, left, Inches(2.75), Inches(1.9), Inches(0.8), name, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, Inches(0.6), Inches(4.4), Inches(12), Inches(2), "演示入口  http://127.0.0.1:5174/    无需登录    使用演示十一幕可带练", size=18, color=MUTED)
    footer(s, 4)

    # 5 architecture
    s = new_light(prs)
    add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "产品整体架构", size=28, bold=True, color=DARK)
    bullets(s, Inches(0.7), Inches(1.2), Inches(12), Inches(5.2), [
        "前端：React + TypeScript + Vite + Ant Design + ECharts",
        "后端：Python + FastAPI + SQLAlchemy + SQLite",
        "AI：XGBoost 预测 · Isolation Forest 异常 · RandomForest 负荷识别 · 规则策略",
        "数据：30 天 × 5 分钟 × 8 类家电，演示日四路对齐 100%",
        "执行：管家宝强模拟控制当前开关与设定温度，不改写历史曲线",
    ], size=20)
    footer(s, 5)

    # 6 features 1
    s = new_light(prs)
    add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "核心功能 · 看懂一户的电", size=28, bold=True, color=DARK)
    bullets(s, Inches(0.7), Inches(1.2), Inches(12), Inches(5.2), [
        "总览：今日约 18.73 kWh，评分、电费、碳排，24 小时回放。",
        "能耗：空调约 41.3%，热水器约 22.3%，冰箱约 11.8%（分路汇总）。",
        "画像与洞察：近 30 日对标，不做虚构的全市平均。",
        "使用说明：五分钟上手，给住在系统里的人看。",
    ], size=20)
    footer(s, 6)

    # 7 features 2
    s = new_light(prs)
    add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "核心功能 · 预测与异常", size=28, bold=True, color=DARK)
    bullets(s, Inches(0.7), Inches(1.2), Inches(12), Inches(5.2), [
        "预测：未来 6/24 小时，绿色带宽 = 真实 RMSE，不是手动画的准确率。",
        "准确率 82.48% 定义为 100−MAPE，不是分类正确率。",
        "异常：03:17 热水器 unusual_time，规则说明时段不合理。",
        "实验对照 13 点、检出 27 条，F1=0.35，页面主动展示。",
    ], size=20)
    footer(s, 7)

    # 8 features 3
    s = new_light(prs)
    add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "核心功能 · 决策到执行", size=28, bold=True, color=DARK)
    bullets(s, Inches(0.7), Inches(1.2), Inches(12), Inches(5.2), [
        "策略：空调 24℃→26℃，舒适滑杆，优化前/后电量、电费、碳排与削峰曲线。",
        "管家宝强：汇报、点头、场景、记忆偏好；冰箱拒绝关闭。",
        "生活系统：电价时钟、今晚任务、电网削峰邀约、谷电预约。",
        "数据来源：四路合同、缺模态降级、接入校验（演示只校验不写库）。",
    ], size=20)
    footer(s, 8)

    # 9 tech
    s = new_light(prs)
    add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "核心技术", size=28, bold=True, color=DARK)
    bullets(s, Inches(0.7), Inches(1.2), Inches(12), Inches(5.2), [
        "16 列预测特征：滞后功率、滚动均值、室内温湿度、占用、小时与星期。不含室外天气。",
        "负荷识别：有分路标签的 1 小时功率窗分类，Accuracy 0.9513。不是总表 NILM。",
        "异常：Isolation Forest 给分，规则引擎给时段和设备。",
        "策略：收益 × 可行性 × 接受度 × 舒适度，可模拟。",
    ], size=20)
    footer(s, 9)

    # 10 flow
    s = new_light(prs)
    add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "技术实现流程", size=28, bold=True, color=DARK)
    bullets(s, Inches(0.7), Inches(1.2), Inches(12), Inches(5.2), [
        "generate_demo_data → init_db → train_models 落盘 joblib 与 metrics.json",
        "FastAPI 读库与模型，返回可复现数字",
        "前端可视化 + 使用演示十一幕 + 管家语音",
        "改 Python 后必须重启 uvicorn；前端 http://127.0.0.1:5174/",
    ], size=20)
    footer(s, 10)

    # 11 innovation
    s = new_light(prs)
    add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "核心创新点", size=28, bold=True, color=DARK)
    bullets(s, Inches(0.7), Inches(1.1), Inches(12), Inches(5.5), [
        "四路时序对齐，缺模态可降级，接入合同可校验。",
        "指标全部来自训练落盘，不拿 LLM 当计算器。",
        "负荷识别边界写清：有标签窗口分类 ≠ 总表盲拆。",
        "异常 F1=0.35 主动展示，用规则补归因。",
        "建议接到管家：能听、能说、能拧、能点头。",
        "特征可解释：占用和室内温度在模型里，室外天气不在。",
    ], size=20)
    footer(s, 11)

    # 12 runtime
    s = new_light(prs)
    add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "产品实际运行效果", size=28, bold=True, color=DARK)
    add_text(s, Inches(0.7), Inches(1.3), Inches(12), Inches(4.8),
             "演示日用电 18.73 kWh\n能源评分按峰值、异常、浪费扣分\n03:17 热水器可回放、可归因\n策略模拟给出削峰曲线\n管家一句话执行今晚方案\n四路对齐 100%",
             size=22, color=TEXT)
    footer(s, 12)

    # 13 experiments
    s = new_light(prs)
    add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "实验 / 测试结果（来自 metrics.json）", size=26, bold=True, color=DARK)
    rows = [
        ["实验", "模型", "关键指标", "说明"],
        ["能耗预测", "XGBoost", "R² 0.9525  MAPE 17.52%  准确率 82.48%", "准确率=100−MAPE"],
        ["异常检测", "IF + 规则", "P 0.26  R 0.54  F1 0.35", "标注 13 / 检出 27"],
        ["负荷识别", "RandomForest", "Acc 0.9513  macro-F1 0.9515", "有标签窗口，非 NILM"],
        ["策略模拟", "规则引擎", "今日 18.73 kWh，空调升温为首选", "实时计算，非口头估算"],
    ]
    table = s.shapes.add_table(len(rows), 4, Inches(0.45), Inches(1.2), Inches(12.4), Inches(4.4)).table
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(2.2)
    table.columns[2].width = Inches(5.0)
    table.columns[3].width = Inches(3.4)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=12, bold=(r == 0), color=DARK if r == 0 else TEXT)
    footer(s, 13)

    # 14 advantage
    s = new_light(prs)
    add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "产品优势", size=28, bold=True, color=DARK)
    bullets(s, Inches(0.7), Inches(1.2), Inches(12), Inches(5.2), [
        "可演示：十一幕使用演示，入口固定，无需登录。",
        "可复现：训练脚本 + metrics.json，数字对得上。",
        "可执行：策略不只停留在报告，管家能拧。",
        "可移植：仿真灌标准表，换适配器和电价即可按户重训。",
        "可诚实：短板（异常 F1、非 NILM）写在产品上。",
    ], size=20)
    footer(s, 14)

    # 15 scenes
    s = new_light(prs)
    add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "应用场景", size=28, bold=True, color=DARK)
    bullets(s, Inches(0.7), Inches(1.2), Inches(12), Inches(5.2), [
        "城市居民家庭日常节能与峰谷错峰。",
        "夜间异常巡检（热水器、待机负荷）。",
        "电网削峰邀约的家庭侧响应（演示）。",
        "能源服务商 / 物业的家庭决策内核（接入真实表后）。",
    ], size=22)
    footer(s, 15)

    # 16 value
    s = new_light(prs)
    add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "产品价值", size=28, bold=True, color=DARK)
    bullets(s, Inches(0.7), Inches(1.2), Inches(12), Inches(5.2), [
        "用户：看懂账单，今晚有动作。",
        "社会：降低高峰与碳排（演示日碳因子 0.525 kg/kWh）。",
        "经济：把可迁移负荷挪到谷电。",
        "技术推广：标准表 + 可降级多模态，便于复制到下一户。",
    ], size=22)
    footer(s, 16)

    # 17 future
    s = new_light(prs)
    add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "产品未来发展", size=28, bold=True, color=DARK)
    bullets(s, Inches(0.7), Inches(1.2), Inches(12), Inches(5.2), [
        "接入真实电表与本地电价配置。",
        "按户重训预测与异常模型。",
        "可选 LLM 解释层，仍不算指标。",
        "物理执行器对接（当前为家庭侧模拟控制）。",
    ], size=22)
    footer(s, 17)

    # 18 team
    s = new_light(prs)
    add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "团队介绍（请填写真实姓名）", size=28, bold=True, color=DARK)
    add_text(s, Inches(0.7), Inches(1.1), Inches(12), Inches(0.5), "本页不得出现指导教师、专家顾问、专利、合同、论文。", size=14, color=GOLD)
    rows = [
        ["成员", "角色", "负责"],
        ["", "产品 / 前端", "页面、交互、使用说明"],
        ["", "后端 / 数据", "API、仿真灌表、生活系统"],
        ["", "AI / 实验", "训练、指标、模型实验舱"],
        ["", "演示 / 材料", "视频、海报、PPT、日志"],
    ]
    table = s.shapes.add_table(len(rows), 3, Inches(0.7), Inches(1.8), Inches(11.8), Inches(4.2)).table
    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(3.4)
    table.columns[2].width = Inches(5.2)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=16, bold=(r == 0), color=DARK)
    footer(s, 18)

    out = MAT / "04_产品说明_智家能.pptx"
    prs.save(out)
    return out


def make_pdf():
    pdfmetrics.registerFont(TTFont("YaHei", str(SIMHEI)))
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="H", fontName="YaHei", fontSize=16, leading=24, textColor=HexColor("#08241c"), spaceAfter=10))
    styles.add(ParagraphStyle(name="B", fontName="YaHei", fontSize=11, leading=18, textColor=HexColor("#123026"), spaceAfter=6))
    styles.add(ParagraphStyle(name="S", fontName="YaHei", fontSize=9, leading=14, textColor=HexColor("#5f746a"), spaceAfter=8))
    md = (MAT / "06_开发日志_智家能.md").read_text(encoding="utf-8")
    story = []
    for raw in md.splitlines():
        line = raw.strip()
        if not line:
            story.append(Spacer(1, 4))
            continue
        if line.startswith("# "):
            story.append(Paragraph(line[2:], styles["H"]))
        elif line.startswith("## "):
            story.append(Paragraph(line[3:], styles["H"]))
        elif line.startswith("### "):
            story.append(Paragraph(line[4:], styles["B"]))
        elif line.startswith("|") or line.startswith("---") or line.startswith("> "):
            story.append(Paragraph(line.replace("|", " ").replace(">", ""), styles["S"]))
        else:
            story.append(Paragraph(line.lstrip("- ").lstrip("* "), styles["B"]))
    out = MAT / "06_开发日志_智家能.pdf"
    doc = SimpleDocTemplate(str(out), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm)
    doc.build(story)
    return out


def copy_proof():
    if PROOF.exists():
        shutil.rmtree(PROOF)
    mapping = {
        "01_产品相关材料": [
            MAT / "01_产品名称.txt",
            MAT / "03_产品介绍_官网填写稿.txt",
            ROOT / "docs" / "产品介绍.md",
            ROOT / "docs" / "创新点.md",
            ROOT / "docs" / "答辩稿.md",
        ],
        "02_系统运行说明": [
            MAT / "05_实物演示视频" / "录制说明与分镜脚本.md",
            ROOT / "docs" / "技术架构.md",
            ROOT / "README.md",
        ],
        "03_实验数据": [
            ROOT / "models" / "metrics.json",
            ROOT / "data" / "processed" / "experiment_results.json",
            ROOT / "docs" / "实验结果.md",
        ],
        "04_测试材料": [
            ROOT / "tests" / "test_api.py",
            ROOT / "tests" / "test_ai.py",
        ],
        "05_项目源代码说明": [],
        "06_产品原型": [
            ROOT / "docs" / "TECHNICAL_SPEC.md",
        ],
        "07_数据集说明": [],
        "08_其他必要证明材料": [
            ROOT / "docs" / "PPT大纲.md",
            MAT / "06_开发日志_智家能.md",
        ],
    }
    for folder, files in mapping.items():
        dest = PROOF / folder
        dest.mkdir(parents=True, exist_ok=True)
        for f in files:
            if f.exists():
                shutil.copy2(f, dest / f.name)

    (PROOF / "05_项目源代码说明" / "说明.txt").write_text(
        "完整工程位于项目根目录，不把 node_modules、.venv、SQLite 大库打进本 ZIP，以免超过 100 MB。\n"
        "根目录含 frontend/ backend/ ai/ scripts/ tests/ models/ data/demo/。\n"
        "启动见 README：前端 http://127.0.0.1:5174/  后端 http://127.0.0.1:8000/docs\n",
        encoding="utf-8",
    )
    (PROOF / "07_数据集说明" / "数据集说明.txt").write_text(
        "演示数据：data/demo/（CSV 灌入 SQLite）。\n"
        "间隔 5 分钟，成都四口之家，夏季，约 30 天。\n"
        "四路：电力 energy_records、室内环境 environment_records、占用行为 behavior_records、天气 weather_records。\n"
        "设备 8 类。异常对照 anomaly_labels。\n"
        "模式为 demo 仿真灌标准表，不是实时电表采集。\n",
        encoding="utf-8",
    )
    (PROOF / "08_其他必要证明材料" / "诚实声明.txt").write_text(
        "预测准确率=100-MAPE。负荷识别不是总表 NILM。异常 F1=0.35 主动展示。\n"
        "对标仅限本户近 30 日。管家控制为模拟执行。\n"
        "Git 当前无提交记录，开发日志已据此如实说明。\n",
        encoding="utf-8",
    )
    poster = MAT / "02_产品海报_智家能_1200x560.png"
    if poster.exists():
        shutil.copy2(poster, PROOF / "01_产品相关材料" / poster.name)

    zip_path = MAT / "07_其他证明材料_智家能.zip"
    if zip_path.exists():
        zip_path.unlink()
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for p in PROOF.rglob("*"):
            if p.is_file():
                zf.write(p, p.relative_to(PROOF.parent))
    return zip_path


def write_checklist(poster: Path, pptx: Path, pdf: Path, zpath: Path):
    def size_mb(p: Path) -> str:
        return f"{p.stat().st_size / 1024 / 1024:.2f} MB" if p.exists() else "缺失"

    text = f"""# 提交清单（对照《材料要求》）

产品名称（证书用）：**智家能**

| 序号 | 材料 | 本目录文件 | 格式/大小 | 状态 |
| -: | --- | --- | --- | --- |
| 1 | 产品名称 | 01_产品名称.txt | 文本 | 已写好 |
| 2 | 产品海报 | {poster.name} | PNG，{size_mb(poster)}，1200×560 | 已生成 |
| 3 | 产品介绍 | 03_产品介绍_官网填写稿.txt | 文本，复制到官网文本框 | 已写好 |
| 4 | 产品说明PPT | {pptx.name} | PPTX，{size_mb(pptx)}，18 页 | 已生成 |
| 5 | 实物演示视频 | 05_实物演示视频/ | 需按分镜自行录制 MP4≤3 分钟 | 待录制 |
| 6 | 开发日志 | {pdf.name} | PDF，{size_mb(pdf)} | 已生成 |
| 7 | 其他证明材料 | {zpath.name} | ZIP，{size_mb(zpath)} | 已打包 |

## 你还需要做的三件事

1. 打开 PPT 最后一页和开发日志第 6 节，填入真实队员姓名。
2. 按 05_实物演示视频/录制说明与分镜脚本.md 录一段不超过 3 分钟的 MP4，放到该文件夹。
3. 上传前再打开一遍 PPT：确认没有指导教师、专利、合同、论文字样。

## 上传时注意

- 海报、PPT、PDF、ZIP、视频分别按比赛系统栏目上传。
- 产品介绍复制 03 号文本，不要贴成本文件。
- 产品名称只填「智家能」。
"""
    (MAT / "00_提交清单与填写说明.md").write_text(text, encoding="utf-8")


def main():
    poster = make_poster()
    pptx = make_pptx(poster)
    pdf = make_pdf()
    zpath = copy_proof()
    write_checklist(poster, pptx, pdf, zpath)
    print("poster", poster, poster.stat().st_size)
    print("pptx", pptx, pptx.stat().st_size)
    print("pdf", pdf, pdf.stat().st_size)
    print("zip", zpath, zpath.stat().st_size)


if __name__ == "__main__":
    main()
